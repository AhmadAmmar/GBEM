"""
Fast header/schema scan for .csv, .shp, .xls, .xlsx (first row only; no full loads).

Usage (PowerShell):
conda run --name eo python "D:/OneDrive - Ulster University/PhD/code/scripts/visualization/fast_header_scan.py" `
  --roots "D:/OneDrive - Ulster University/PhD/data" `
  --skip "all-domestic-certificates" "EnglandWalesSheets" `
  --out  "D:/OneDrive - Ulster University/PhD/Maps/_inventory"

Outputs:
- inventory_headers_YYYYMMDD_HHMMSS.csv  (path, kind, sheet, columns_first_row, schema_props/crs for SHP)
"""

from __future__ import annotations
import os, csv as pycsv, argparse, json
from datetime import datetime

# Optional libs (we handle missing gracefully)
try:
    import fiona
    HAVE_FIONA = True
except Exception:
    HAVE_FIONA = False

try:
    import openpyxl  # for .xlsx
    HAVE_OPENPYXL = True
except Exception:
    HAVE_OPENPYXL = False

try:
    import xlrd  # for legacy .xls
    HAVE_XLRD = True
except Exception:
    HAVE_XLRD = False


# ---------------- Config defaults ----------------
DEFAULT_ROOTS = [r"D:/OneDrive - Ulster University/PhD/data"]
DEFAULT_SKIP  = ["all-domestic-certificates", "EnglandWalesSheets", "~$"]
SCAN_EXTS     = {".csv", ".shp", ".xls", ".xlsx"}
# -------------------------------------------------


def norm(p: str) -> str:
    return os.path.normpath(os.path.abspath(p))


def should_skip_dir(dirpath: str, skip_substrings: list[str]) -> bool:
    d = dirpath.lower()
    return any(s.lower() in d for s in skip_substrings)


def classify(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":  return "csv"
    if ext == ".shp":  return "shp"
    if ext == ".xlsx": return "xlsx"
    if ext == ".xls":  return "xls"
    return "other"


def read_csv_header(path: str) -> dict:
    # Only read the first non-empty line as header
    try:
        with open(path, "r", encoding="utf-8", errors="replace", newline="") as f:
            reader = pycsv.reader(f)
            for row in reader:
                if row and any(cell.strip() for cell in row):
                    return {"columns_first_row": row}
        return {"columns_first_row": []}
    except Exception as e:
        return {"error": f"{type(e).__name__}: {e}"}


def read_xlsx_first_row(path: str) -> dict:
    if not HAVE_OPENPYXL:
        return {"error": "openpyxl not installed"}
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = wb.sheetnames[0] if wb.sheetnames else None
        cols = []
        if sheet:
            ws = wb[sheet]
            # first row only
            row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
            if row:
                cols = ["" if c is None else str(c) for c in row]
        wb.close()
        return {"sheet": sheet, "columns_first_row": cols}
    except Exception as e:
        return {"error": f"{type(e).__name__}: {e}"}


def read_xls_first_row(path: str) -> dict:
    if not HAVE_XLRD:
        return {"error": "xlrd not installed"}
    try:
        book = xlrd.open_workbook(path, on_demand=True)
        sheet = book.sheet_names()[0] if book.nsheets else None
        cols = []
        if sheet:
            sh = book.sheet_by_index(0)
            # first row only
            cols = [str(v) for v in sh.row_values(0)]
        book.release_resources()
        return {"sheet": sheet, "columns_first_row": cols}
    except Exception as e:
        return {"error": f"{type(e).__name__}: {e}"}


def read_shp_schema(path: str) -> dict:
    if not HAVE_FIONA:
        return {"error": "fiona not installed"}
    try:
        with fiona.open(path) as src:
            props = list(src.schema.get("properties", {}).keys())
            geom  = src.schema.get("geometry")
            crs   = src.crs_wkt or src.crs
        return {"geometry": geom, "crs": str(crs) if crs else None, "schema_props": props}
    except Exception as e:
        return {"error": f"{type(e).__name__}: {e}"}


def scan_file(path: str) -> dict:
    kind = classify(path)
    info = {"path": path, "kind": kind}
    if kind == "csv":
        info.update(read_csv_header(path))
    elif kind == "xlsx":
        info.update(read_xlsx_first_row(path))
    elif kind == "xls":
        info.update(read_xls_first_row(path))
    elif kind == "shp":
        info.update(read_shp_schema(path))
    else:
        info["note"] = "skipped"
    return info


def write_csv(rows: list[dict], out_dir: str) -> str:
    os.makedirs(out_dir, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out = os.path.join(out_dir, f"inventory_headers_{ts}.csv")
    fields = ["path", "kind", "sheet", "columns_first_row", "geometry", "crs", "schema_props", "error", "note"]
    with open(out, "w", newline="", encoding="utf-8") as f:
        w = pycsv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for r in rows:
            row = r.copy()
            for k in ("columns_first_row", "schema_props"):
                if k in row and isinstance(row[k], list):
                    row[k] = json.dumps(row[k], ensure_ascii=False)
            w.writerow({k: row.get(k) for k in fields})
    return out


def main():
    ap = argparse.ArgumentParser(description="Fast first-row/schema scan for csv/shp/xls/xlsx.")
    ap.add_argument("--roots", nargs="+", default=DEFAULT_ROOTS, help="Root directory/ies to scan.")
    ap.add_argument("--skip",  nargs="*",  default=DEFAULT_SKIP,  help="Skip any directory whose path contains these substrings.")
    ap.add_argument("--out",   default=r"D:/OneDrive - Ulster University/PhD/Maps/_inventory", help="Output directory for CSV report.")
    args = ap.parse_args()

    roots = [norm(r) for r in args.roots]
    skip  = [s.lower() for s in args.skip]

    print("\n=== FAST HEADER SCAN ===")
    print("Roots:", *roots, sep="\n - ")
    print("Skip directory substrings:", skip)
    print("Extensions: .csv .shp .xls .xlsx\n")

    rows: list[dict] = []
    for root in roots:
        for dirpath, dirnames, filenames in os.walk(root):
            if should_skip_dir(dirpath, skip):
                dirnames[:] = []  # prune
                continue
            for fn in filenames:
                ext = os.path.splitext(fn)[1].lower()
                if ext not in SCAN_EXTS:
                    continue
                path = os.path.join(dirpath, fn)
                info = scan_file(path)
                rows.append(info)

                # concise console line
                base = path.replace(root, "").lstrip("\\/")
                if "error" in info and info["error"]:
                    print(f"[{info['kind']:5}] {base}\n   ↳ ERROR: {info['error']}")
                elif info["kind"] == "shp":
                    props = info.get("schema_props") or []
                    preview = ", ".join(props[:12]) + (" …" if len(props) > 12 else "")
                    print(f"[shp  ] {base}\n   ↳ geom={info.get('geometry')} crs={info.get('crs')} | props: {preview}")
                else:
                    cols = info.get("columns_first_row") or []
                    preview = ", ".join(map(str, cols[:12])) + (" …" if len(cols) > 12 else "")
                    print(f"[{info['kind']:5}] {base}\n   ↳ {preview}")

    out_csv = write_csv(rows, norm(args.out))
    print(f"\nSaved report: {out_csv}\nDone.\n")


if __name__ == "__main__":
    main()
